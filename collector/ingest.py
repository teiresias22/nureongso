#!/usr/bin/env python3
"""누렁소검은소 수집기.

사용:
    python ingest.py members            # 역대+현역 의원
    python ingest.py bills --age 22     # 발의법률안 + 발의/공동발의
    python ingest.py plenary --age 22   # 본회의 처리 의안
    python ingest.py votes --age 22     # 본회의 표결 (plenary 먼저 실행)
    python ingest.py summaries --age 22 # 법안 제안이유·주요내용
    python ingest.py sidejobs           # 겸직 결정 내역 (20대~, 통째로 교체)
    python ingest.py attendance --age 22 # 본회의 출결 누적 (최신 회기 엑셀 하나)
    python ingest.py trips              # 직무상 국외활동 신고 (21대~, 통째로 교체)
    python ingest.py research           # 의원 연구단체 (20~22대, 통째로 교체)
    python ingest.py refresh            # member_stats 갱신
    python ingest.py all --age 22

환경변수: DATABASE_URL (Supabase > Settings > Database > Connection string),
          ASSEMBLY_API_KEY (필수. https://open.assembly.go.kr 무료 발급)
"""
from __future__ import annotations

import argparse
import html
import io
import json
import os
import re
import sys
import time
import unicodedata

import httpx
import psycopg

BASE = "https://open.assembly.go.kr/portal/openapi"
KEY = os.getenv("ASSEMBLY_API_KEY") or None
PAGE = 1000  # API 최대

SERVICES = {
    "incumbent": "nwvrqwxyaytdsfvhu",  # 현역 국회의원 인적사항
    "allmember": "ALLNAMEMBER",  # 역대 국회의원 인적사항
    "bills": "nzmimeepazxkubdpn",  # 국회의원 발의법률안
    "plenary": "ncocpgfiaoituanbr",  # 본회의 처리 의안 (표결 집계)
    "votes": "nojepdqqaweusdfbi",  # 국회의원 본회의 표결
    "summary": "BPMBILLSUMMARY",  # 법률안 제안이유 및 주요내용 (BILL_NO 필수)
    # 목록(OPENSRVAPI)의 SRV_URL 은 설명 화면 주소다. 호출 이름은 명세서 xls 에만 있다.
    "sidejob": "nahfbzwvatmaxscwq",  # 국회의원 겸직 결정 내역 (OHAC6C000892WC13765)
    "trips": "nasnutdbapnfphwyr",  # 국회의원 직무상 국외활동 신고 내역 (O87UNV000897E818234)
    "research": "numwhtqhavaqssfle",  # 국회의원 연구단체 등록현황 (O78HKE0010099W15881, REGDAESU 필수)
}


def fetch(service: str, max_rows: int | None = None, **params) -> list[dict]:
    """서비스 전체 row 를 페이지 순회로 가져온다. max_rows 로 조기 종료(테스트용).

    인증키 없이 호출하면 서버가 pIndex/pSize 를 무시하고 매번 첫 5행만 돌려준다.
    (총건수 head 는 정상값을 주므로 조용히 같은 행을 반복 수집하게 된다.)
    그래서 KEY 는 필수이고, 그래도 페이지가 안 넘어가면 아래에서 중단한다.
    """
    if not KEY:
        raise RuntimeError(
            "ASSEMBLY_API_KEY 가 필요합니다. 키 없이 호출하면 서버가 페이지네이션을 무시하고"
            " 첫 5행만 반복해서 돌려줍니다. https://open.assembly.go.kr 에서 무료 발급하세요."
        )
    name = SERVICES[service]
    out: list[dict] = []
    seen: set[str] = set()
    page = 1
    with httpx.Client(timeout=60, headers={"User-Agent": "nureongso/0.1"}) as c:
        while True:
            q = {"Type": "json", "pIndex": page, "pSize": PAGE, "KEY": KEY, **params}
            # 국회 서버가 간헐적으로 TCP 연결 자체를 안 받는다(connect timeout).
            # 하루 한 번 도는 작업이라 몇 분 더 기다리는 편이 하루를 건너뛰는 것보다 낫다.
            for attempt in range(5):
                try:
                    data = c.get(f"{BASE}/{name}", params=q).json()
                    break
                except Exception as e:  # 네트워크/JSON 오류만 재시도
                    if attempt == 4:
                        raise
                    wait = 5 * 2**attempt  # 5, 10, 20, 40초
                    print(f"  retry {attempt + 1} ({wait}초 뒤): {e}", file=sys.stderr)
                    time.sleep(wait)

            if "RESULT" in data:  # 에러 또는 데이터 없음
                code = data["RESULT"]["CODE"]
                if code.startswith("INFO-200"):  # 해당 데이터 없음
                    break
                raise RuntimeError(f"{name}: {data['RESULT']}")

            body = data[name]
            rows = next((b["row"] for b in body if "row" in b), [])
            total = body[0]["head"][0]["list_total_count"]

            # 페이지가 실제로 넘어갔는지 확인. 서버가 pIndex 를 무시하면 같은 행이 다시 온다.
            fresh = [r for r in rows if (k := json.dumps(r, sort_keys=True)) not in seen and not seen.add(k)]
            if rows and not fresh:
                raise RuntimeError(
                    f"{name}: pIndex={page} 가 이전 페이지와 같은 행을 반환했습니다."
                    " 인증키가 유효한지 확인하세요."
                )
            out.extend(fresh)
            print(f"  {name} {len(out)}/{total}", file=sys.stderr)
            if len(out) >= total or not rows or (max_rows and len(out) >= max_rows):
                break
            page += 1
    return out


def upsert(cur, table: str, cols: list[str], rows: list[tuple], conflict: str, update: bool = True):
    if not rows:
        return
    ph = ",".join(["%s"] * len(cols))
    collist = ",".join(cols)
    if update:
        sets = ",".join(f"{c}=excluded.{c}" for c in cols if c not in conflict.split(","))
        action = f"do update set {sets}" if sets else "do nothing"
    else:
        action = "do nothing"
    sql = f"insert into {table} ({collist}) values ({ph}) on conflict ({conflict}) {action}"
    cur.executemany(sql, rows)


def d(v):
    """빈 문자열 -> None, 날짜 문자열 정리."""
    if v is None:
        return None
    v = str(v).strip()
    return v or None


# --------------------------------------------------------------------------- 의원


def ingest_members(cur) -> int:
    hist = fetch("allmember")
    inc = fetch("incumbent")
    incumbent_codes = {r["MONA_CD"] for r in inc if r.get("MONA_CD")}

    rows = []
    for r in hist:
        code = d(r.get("NAAS_CD"))
        if not code:
            continue
        rows.append(
            (
                code,
                d(r.get("NAAS_NM")),
                d(r.get("NAAS_CH_NM")),
                d(r.get("BIRDY_DT")),
                d(r.get("NTR_DIV")),
                d(r.get("PLPT_NM")),
                d(r.get("ELECD_NM")),
                d(r.get("ELECD_DIV_NM")),
                d(r.get("GTELT_ERACO")),
                d(r.get("RLCT_DIV_NM")),
                d(r.get("BLNG_CMIT_NM")),
                d(r.get("NAAS_PIC")),
                d(r.get("NAAS_TEL_NO")),
                d(r.get("NAAS_EMAIL_ADDR")),
                d(r.get("NAAS_HP_URL")),
                code in incumbent_codes,
            )
        )
    cols = [
        "code", "name", "name_hanja", "birth", "sex", "party", "district",
        "elect_type", "terms", "term_count", "committees", "photo_url",
        "tel", "email", "homepage", "is_incumbent",
    ]
    # member 에는 국회의원만 있는 게 아니다. 이 API 는 역대 국회의원을 주므로 국회의원
    # 출신 단체장·교육감도 여기 걸린다. 그 사람들의 값을 덮으면 지금 직위가 아니라
    # 의원 시절로 되돌아간다 — 오세훈 서울시장이 '한나라당 · 강남구을' 이 된다.
    # 이건 매일 도는 수집이라 손으로 고쳐도 다음 날 아침이면 제자리다.
    #
    # 그래서 사람 자체를 가리키는 값만 덮고, '지금 무슨 자리에 있나' 에 달린 값은
    # 국회의원인 사람에게만 덮는다. 다른 직위는 선관위 수집기(nec.py)가 맡는다.
    IDENTITY = {"name", "name_hanja", "birth", "sex"}
    placeholders = ",".join(["%s"] * len(cols))
    sets = ",".join(
        f"{c}=excluded.{c}" if c in IDENTITY else
        f"{c}=case when coalesce(member.office,'국회의원') = '국회의원'"
        f"         then excluded.{c} else member.{c} end"
        for c in cols if c != "code"
    )
    cur.executemany(
        f"insert into member ({','.join(cols)}) values ({placeholders})"
        f" on conflict (code) do update set {sets}",
        rows,
    )

    # 현역은 최신값(정당/지역구/위원회)이 현역 API 쪽이 정확하므로 덮어쓴다.
    inc_rows = [
        (
            d(r.get("POLY_NM")), d(r.get("ORIG_NM")), d(r.get("ELECT_GBN_NM")),
            d(r.get("CMIT_NM")), d(r.get("REELE_GBN_NM")), d(r.get("UNITS")),
            d(r.get("TEL_NO")), d(r.get("E_MAIL")), d(r.get("HOMEPAGE")),
            d(r.get("MONA_CD")),
        )
        for r in inc if r.get("MONA_CD")
    ]
    cur.executemany(
        "update member set party=%s, district=%s, elect_type=%s, committees=%s,"
        " term_count=%s, terms=%s, tel=%s, email=%s, homepage=%s,"
        " is_incumbent=true, office='국회의원'"
        " where code=%s",
        inc_rows,
    )
    return len(rows)


# --------------------------------------------------------------------------- 법안


def ingest_bills(cur, age: int) -> int:
    rows = fetch("bills", AGE=age)
    bills, sponsors = [], []
    for r in rows:
        bid = d(r.get("BILL_ID"))
        if not bid:
            continue
        bills.append(
            (
                bid, age, d(r.get("BILL_NO")), d(r.get("BILL_NAME")),
                d(r.get("COMMITTEE")), d(r.get("PROPOSE_DT")),
                d(r.get("PROC_RESULT")), d(r.get("PROC_DT")),
                d(r.get("PROPOSER")), d(r.get("DETAIL_LINK")),
            )
        )
        if rst := d(r.get("RST_MONA_CD")):
            for code in rst.split(","):
                if code.strip():
                    sponsors.append((bid, code.strip(), "rep"))
        if publ := d(r.get("PUBL_MONA_CD")):
            for code in publ.split(","):
                if code.strip():
                    sponsors.append((bid, code.strip(), "co"))

    upsert(
        cur, "bill",
        ["bill_id", "age", "bill_no", "name", "committee", "proposed_at",
         "proc_result", "proc_dt", "proposer", "detail_link"],
        bills, "bill_id",
    )
    # 대표발의자가 공동발의 명단에도 들어간 경우 중복 제거
    sponsors = list(dict.fromkeys(sponsors))
    upsert(cur, "bill_sponsor", ["bill_id", "member_code", "role"],
           sponsors, "bill_id,member_code,role", update=False)
    print(f"  bills={len(bills)} sponsors={len(sponsors)}", file=sys.stderr)
    return len(bills)


def ingest_plenary(cur, age: int) -> int:
    rows = fetch("plenary", AGE=age)
    out = [
        (
            d(r.get("BILL_ID")), age, d(r.get("BILL_NO")), d(r.get("BILL_NAME")),
            d(r.get("CURR_COMMITTEE")), d(r.get("PROC_DT")), d(r.get("PROC_RESULT_CD")),
            r.get("YES_TCNT"), r.get("NO_TCNT"), r.get("BLANK_TCNT"),
            r.get("VOTE_TCNT"), r.get("MEMBER_TCNT"),
        )
        for r in rows if r.get("BILL_ID")
    ]
    upsert(
        cur, "plenary_bill",
        ["bill_id", "age", "bill_no", "name", "committee", "proc_dt", "proc_result",
         "yes_cnt", "no_cnt", "blank_cnt", "vote_cnt", "member_cnt"],
        out, "bill_id",
    )
    return len(out)


def ingest_votes(cur, age: int, only_new: bool = True) -> int:
    """표결은 의안 단위 호출이라 건수가 많다. 이미 받은 의안은 건너뛴다."""
    cur.execute("select bill_id from plenary_bill where age = %s", (age,))
    todo = [r[0] for r in cur.fetchall()]
    if only_new:
        cur.execute("select distinct bill_id from vote")
        done = {r[0] for r in cur.fetchall()}
        todo = [b for b in todo if b not in done]

    total = 0
    for i, bid in enumerate(todo, 1):
        rows = fetch("votes", AGE=age, BILL_ID=bid)
        out = [
            (bid, d(r.get("MONA_CD")), d(r.get("RESULT_VOTE_MOD")), parse_dt(r.get("VOTE_DATE")))
            for r in rows if r.get("MONA_CD")
        ]
        upsert(cur, "vote", ["bill_id", "member_code", "result", "voted_at"],
               out, "bill_id,member_code")
        total += len(out)
        if i % 20 == 0:
            print(f"  votes {i}/{len(todo)} bills, {total} rows", file=sys.stderr)
    return total


def parse_dt(v):
    """'20260917 161700' -> ISO."""
    v = d(v)
    if not v or len(v) < 8:
        return None
    date, _, t = v.partition(" ")
    t = (t + "000000")[:6]
    return f"{date[:4]}-{date[4:6]}-{date[6:8]} {t[:2]}:{t[2:4]}:{t[4:6]}"



def ingest_summaries(cur, age: int, limit: int | None = None) -> int:
    """법안 제안이유·주요내용. 의안 1건당 1회 호출이라 없는 것만 받는다.

    BPMBILLSUMMARY 는 BILL_ID 를 안 받고 BILL_NO 만 받는다. AGE 만 주면 ERROR-300 이다.
    """
    cur.execute(
        "select bill_id, bill_no from bill"
        " where age = %s and bill_no is not null and summary is null"
        " order by proposed_at desc", (age,))
    todo = cur.fetchall()
    if limit:
        todo = todo[:limit]
    print(f"  요약 없는 의안 {len(todo)}건", file=sys.stderr)

    got = 0
    for i, (bid, bno) in enumerate(todo, 1):
        try:
            rows = fetch("summary", BILL_NO=bno)
        except Exception as e:
            print(f"  [{i}/{len(todo)}] {bno} 실패: {str(e)[:70]}", file=sys.stderr)
            continue
        # BILL_NO 는 대수가 다르면 겹칠 수 있어 BILL_ID 가 같은 행을 고른다.
        hit = next((r for r in rows if d(r.get("BILL_ID")) == bid), None)
        text = d(hit.get("SUMMARY")) if hit else None
        if text:
            cur.execute("update bill set summary = %s where bill_id = %s",
                        (text.replace("\x00", ""), bid))
            got += 1
        if i % 50 == 0 or i == len(todo):
            cur.connection.commit()  # 오래 걸리는 단계라 부분 결과를 남긴다
            print(f"  [{i}/{len(todo)}] 수집 {got}건", file=sys.stderr)
    return got


# --------------------------------------------------------------------------- 이름으로 잇기

HANJA = re.compile(r"^[一-鿿]+$")


def people_of(cur, age: int) -> list[tuple]:
    """그 대수에 재직한 사람. (code, name, name_hanja, district, elect_type, party, 지역구당선).

    지역구당선 = 그 대수 총선~임기 끝 사이에 지역구(sg 2)로 당선된 기록이 있는가.
    member.district·elect_type 은 최신 값만 있어서 21대 비례였다가 22대 지역구로 옮긴
    이수진(성남중원)은 '비례' 로 안 잡힌다. '(비)' 단서는 이 칸으로 가른다.

    name_hanja 는 NFC 로 편다. 열린국회정보가 준 한자가 호환용 코드로 들어 있다 — 李 가
    U+674E 가 아니라 U+F9E1 이다(실측 790명). 겸직 API 는 표준 코드로 보내서 그대로
    비교하면 같은 글자가 안 맞는다.
    """
    y = 2016 + 4 * (age - 20)  # 그 대수를 뽑은 총선 해
    cur.execute(
        "select code, name, name_hanja, district, elect_type, party,"
        " exists (select 1 from candidacy c where c.member_code = m.code and c.elected"
        "   and c.sg_typecode = '2' and c.election_id between %s and %s)"
        " from member m where terms like %s",
        # 끝은 다음 총선 해 1월 1일. 5월 29일까지 잡으면 다음 총선(4월)이 들어온다.
        (f"{y}0101", f"{y + 4}0101", f"%제{age}대%"))
    return [(c, n, unicodedata.normalize("NFC", h) if h else h, *rest)
            for c, n, h, *rest in cur.fetchall()]


# 개명·합당 계보. member.party 는 수집 당시 이름으로 굳어 있고('미래통합당'), 출결표는
# 그 회기 이름('국민의힘')으로 적는다. 21대 김병욱 둘은 한자까지 같아 정당으로만 갈린다.
PARTY_LINE = {
    "자유한국당": "국민의힘", "미래통합당": "국민의힘", "미래한국당": "국민의힘", "국민의미래": "국민의힘",
    "더불어시민당": "더불어민주당", "더불어민주연합": "더불어민주당",
    # 20대 최경환(광주 북구을)은 국민의당→민주평화당→대안신당. 연구단체 명단이 마지막 당명으로 적는다.
    "민주평화당": "국민의당", "대안신당": "국민의당",
}


def same_party(member_party: str | None, party: str) -> bool:
    """member.party 는 'A/B' 이력일 수 있다. 계보로 펴서 하나라도 같으면 같은 당."""
    norm = lambda v: PARTY_LINE.get(v.strip(), v.strip())
    return norm(party) in {norm(v) for v in (member_party or "").split("/") if v.strip()}


def match_person(people: list[tuple], raw: str, party: str | None = None,
                 used: set[str] = frozenset()) -> str | None:
    """의원 코드 없이 이름만 주는 자료(겸직·출결)를 사람에 잇는다.

    실측으로 본 이름 모양: '박 정'(띄어쓰기), '李達坤'(한자만), '최경환(崔敬煥)'(한자 병기),
    '이수진(비)'(비례 표시). 동명이인은 괄호 → 정당 → 이미 붙은 사람 제외 순으로 가른다.
    출결표의 22대 박지원 둘은 한 사람이 '朴芝源' 으로 적혀 있어 한자 쪽을 먼저 붙이면
    나머지 '박지원' 이 저절로 갈린다. 그래도 하나가 아니면 비운다.
    """
    raw = unicodedata.normalize("NFC", raw.strip())
    m = re.match(r"^(.+?)\((.+)\)$", raw)
    base, hint = (m.group(1), m.group(2)) if m else (raw, None)
    base = base.replace(" ", "")
    if HANJA.match(base):
        hits = [p for p in people if p[2] == base]
    else:
        hits = [p for p in people if p[1] == base]
    if len(hits) > 1 and hint:
        if HANJA.match(hint):
            hits = [p for p in hits if p[2] == hint]
        elif hint.startswith("비"):
            hits = ([p for p in hits if "비례" in f"{p[3]}{p[4]}"]
                    or [p for p in hits if len(p) > 6 and not p[6]])
        else:
            hits = [p for p in hits if hint.split()[-1] in (p[3] or "")]
    if len(hits) > 1 and party:
        hits = [p for p in hits if same_party(p[5], party)] or hits
    if len(hits) > 1 and used:
        hits = [p for p in hits if p[0] not in used]
    return hits[0][0] if len(hits) == 1 else None


# --------------------------------------------------------------------------- 겸직


def ymd(v) -> str | None:
    """'2021.2.22.' / '2024-09-20' → '2021-02-22'. 한 API 안에 두 형식이 섞여 온다."""
    nums_ = re.findall(r"\d+", str(v or ""))
    if len(nums_) < 3:
        return None
    y, mo, da = (int(x) for x in nums_[:3])
    return f"{y:04d}-{mo:02d}-{da:02d}"


def sidejob_kind(decision: str | None) -> str:
    """결정 내용 원문은 열두 가지로 적혀 온다. 화면에서 가를 것은 셋뿐이다."""
    s = (decision or "").replace(" ", "")
    if "사직" in s:
        return "사직권고"
    if "불가" in s or "해당하지않음" in s:
        return "불가"
    return "허용"


def ingest_sidejobs(cur, age: int) -> int:
    """겸직 결정 내역. 20~22대 수백 건이라 매번 통째로 바꾼다.

    국회법 제29조: 의원이 다른 직을 가지면 신고하고, 의장이 윤리심사자문위원회 의견을
    들어 허용 여부를 정한다. 사직권고·겸직 불가도 그대로 공개된다.
    """
    rows = fetch("sidejob")
    ages = {int(re.sub(r"\D", "", r.get("ORD_NUM") or "") or 0) for r in rows} - {0}
    people = {a: people_of(cur, a) for a in ages}
    out, miss = [], []
    for r in rows:
        a = int(re.sub(r"\D", "", r.get("ORD_NUM") or "") or 0)
        name = d(r.get("PN")) or ""
        code = match_person(people.get(a, []), name)
        if not code:
            miss.append(f"{name}({a}대)")
        out.append((a, d(r.get("YR")), ymd(r.get("OPB_DAY")), name, code,
                    d(r.get("CCOF_INST_NM")), d(r.get("PSIT_NM")),
                    d(r.get("CCOF_PSB_YN_CD")), sidejob_kind(r.get("CCOF_PSB_YN_CD"))))
    cur.execute("delete from member_sidejob")
    cur.executemany(
        "insert into member_sidejob (age, year, opened_at, name, member_code, org, position,"
        " decision, decision_kind) values (%s,%s,%s,%s,%s,%s,%s,%s,%s)", out)
    if miss:
        print(f"  ! 사람을 못 찾은 줄 {len(miss)}: {', '.join(miss[:20])}", file=sys.stderr)
    return len(out)


# --------------------------------------------------------------------------- 국외활동


def trip_period(v: str | None) -> tuple[str | None, str | None]:
    """'2021. 4.11~ 4.13' / '2022.12.28.~1.3.' / '2023.5.22.~5.26.(25)' → (시작일, 끝일).

    끝은 월·일만 적는 게 보통이고 해를 넘기면 월이 작아진다. 괄호 안 숫자는 버린다.
    의원마다 돌아온 날이 다르면 '8.17./16.', '3.1/2/3.' 처럼 빗금으로 잇는다 — 첫 날짜를 쓴다.
    """
    head, _, tail = re.sub(r"\(.*?\)", "", v or "").partition("~")
    tail = tail.split("/")[0]
    a = [int(x) for x in re.findall(r"\d+", head)]
    b = [int(x) for x in re.findall(r"\d+", tail)]
    if len(a) < 3:
        return None, None
    y, mo, da = a[:3]
    start = f"{y:04d}-{mo:02d}-{da:02d}"
    if len(b) >= 3:
        ey, em, ed = b[:3]
    elif len(b) == 2:
        em, ed = b
        ey = y + 1 if em < mo else y
    else:
        return start, start
    return start, f"{ey:04d}-{em:02d}-{ed:02d}"


# 원문 오타. 함께 간 사람이 이준석(개혁신당)이고 그 대수에 같은 이름의 의원이 없다 —
# 개혁신당 의원 이주영·천하람이다. 새 오타는 '사람을 못 찾은 줄' 로 드러난다.
TRIP_NAME_FIX = {(22, "이준영"): "이주영", (22, "천아람"): "천하람"}


def ingest_trips(cur, age: int) -> int:
    """직무상 국외활동 신고(국회의원윤리실천규범). 21·22대 합쳐 200건 남짓이라 통째로 바꾼다.
    경비를 누가 댔는지(자비·외교부·외국 정부·민간단체)와 결과보고서 제출 여부가 함께 온다.

    한 줄에 함께 간 의원 이름이 쉼표로 나열돼 온다(' 김원이, 김정호, 박지혜'). 사람마다 펴고,
    원문 목록은 companions 에 그대로 둔다. 경비 기관은 원문이 제각각이라('자비', '자부담',
    '외교부(국가기관)', 줄바꿈 뒤 주석까지) 가르지 않고 공백만 정리해 싣는다.
    """
    rows = fetch("trips")
    people = {a: people_of(cur, a) for a in {int(re.sub(r"\D", "", r.get("UNIT_CD") or "")[-2:] or 0)
                                             for r in rows} - {0}}
    clean = lambda v: re.sub(r"\s+", " ", html.unescape(v or "")).strip() or None
    out, miss = [], []
    for r in rows:
        a = int(re.sub(r"\D", "", r.get("UNIT_CD") or "")[-2:] or 0)
        names = [n.strip() for n in re.split(r"[,，、]", r.get("PN") or "") if n.strip()]
        start, end = trip_period(r.get("SCH_DYS"))
        for name in names:
            code = match_person(people.get(a, []), TRIP_NAME_FIX.get((a, name), name))
            if not code:
                miss.append(f"{name}({a}대)")
            out.append((a, name, code, ", ".join(names), clean(r.get("DSTN_NM")),
                        clean(r.get("PURP_RSON")), clean(r.get("SCH_DYS")), start, end,
                        clean(r.get("EXPNS_SUPPT_INST_NM")), (r.get("REPORT_YN") or "").strip() == "유"))
    cur.execute("delete from member_trip")
    cur.executemany(
        "insert into member_trip (age, name, member_code, companions, destination, purpose, period,"
        " start_on, end_on, funder, reported) values (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)", out)
    if miss:
        print(f"  ! 사람을 못 찾은 줄 {len(miss)}: {', '.join(miss[:20])}", file=sys.stderr)
    return len(out)


# --------------------------------------------------------------------------- 연구단체

RESEARCH_AGES = (20, 21, 22)
RESEARCH_ROLES = (("MAIN_MEM", "대표"), ("RE_MEM", "연구책임"), ("OBJ_MEM", "구성"))


def research_people(v: str | None) -> list[tuple[str, str]]:
    """'김기현(국민의힘), 이준석(개혁신당)' → [('김기현', '국민의힘'), ('이준석', '개혁신당')].
    대표·연구책임도 공동이면 쉼표로 여럿 온다."""
    return [(n.strip(), p.strip()) for n, p in re.findall(r"([^,()]+)\(([^()]+)\)", v or "")]


def ingest_research(cur, age: int) -> int:
    """의원 연구단체. 대수마다 70개 남짓이라 20~22대를 통째로 바꾼다.

    구성인원('21명 : …')이 대표 + 연구책임 + 구성의원 수와 같다 — 구성의원 목록에 대표와
    연구책임은 들어 있지 않다(실측). 이름에 정당이 붙어 와서 동명이인은 정당으로 가른다.
    """
    out, miss = [], []
    for a in RESEARCH_AGES:
        people = people_of(cur, a)
        # 원문에 '&lsquo;', '&middot;' 같은 HTML 문자 참조가 섞여 온다.
        clean = lambda v: re.sub(r"\s+", " ", html.unescape(v or "")).strip() or None
        for r in fetch("research", REGDAESU=str(a)):
            for col, role in RESEARCH_ROLES:
                for name, party in research_people(r.get(col)):
                    # 20대 김성태 둘은 '김성태-지'·'김성태-비' 로 적혀 온다(지역구·비례).
                    base, tag = (name[:-2], name[-1]) if re.search(r"-[지비]$", name) else (name, None)
                    pool = [p for p in people if tag is None or p[6] == (tag == "지")]
                    code = match_person(pool, base, party)
                    if not code:
                        miss.append(f"{name}({a}대)")
                    out.append((a, clean(r.get("RE_NAME")), clean(r.get("RE_TOPIC_NAME")),
                                clean(r.get("RE_OBJECTIVE")), role, name, party, code,
                                clean(r.get("MEMBER_CNT")), clean(r.get("LINK_URL"))))
    cur.execute("delete from member_research")
    cur.executemany(
        "insert into member_research (age, group_name, topic, objective, role, name, party,"
        " member_code, member_cnt, link_url) values (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)", out)
    if miss:
        print(f"  ! 사람을 못 찾은 줄 {len(miss)}: {', '.join(miss[:20])}", file=sys.stderr)
    return len(out)


# --------------------------------------------------------------------------- 출결

# 본회의 출결은 Open API 가 없다. 열린국회정보 '파일 데이터' 로 회기마다 엑셀 하나가
# 올라온다(22대 제415회부터 xlsx, 그 앞은 pdf). 최신 파일에 그 대수 **누적 총계** 가
# 있어서 하나만 받으면 된다 — 438회와 437회 파일을 대조하니 전원이 그 회기 회의일수만큼
# 정확히 늘었다. 공개 API 가 아니라 화면이 쓰는 엔드포인트라 모양이 바뀌면 멈춘다.
ATT_INF = "O4Q5B50011905O18367"
ATT_LIST = "https://open.assembly.go.kr/portal/data/file/searchFileData.do"
ATT_DOWN = "https://open.assembly.go.kr/portal/data/file/downloadFileData.do"
ATT_COLS = ("회의일수", "출석", "결석", "청가", "출장", "결석신고서")


ATT_PAGE = f"https://open.assembly.go.kr/portal/data/service/selectServicePage.do?infId={ATT_INF}&infSeq=1"


def attendance_files(c: httpx.Client) -> list[dict]:
    """출결 파일 목록. 회기 번호(no)를 붙여 오래된 순으로."""
    r = c.post(ATT_LIST, data={"infId": ATT_INF, "infSeq": 1},
               headers={"X-Requested-With": "XMLHttpRequest"})
    r.raise_for_status()
    files = (r.json() or {}).get("data")
    if not files:
        raise RuntimeError("출결 파일 목록 형태가 바뀌었습니다")
    out = []
    for f in files:
        m = re.search(r"제(\d+)회", f.get("viewFileNm", ""))
        if m:
            out.append({**f, "no": int(m.group(1))})
    return sorted(out, key=lambda f: f["no"])


def age_at(day: str) -> int:
    """그 날짜의 국회 대수 (asset.py 의 age_at 과 같은 규칙: 제20대가 2016-05-30 시작)."""
    y, m, dd = (int(x) for x in day.split("-"))
    return 20 + (y - 2016 - (1 if (m, dd) < (5, 30) else 0)) // 4


ATT_LINE = re.compile(r"^(\S+)\s+(\S+)\s+(.*?)\s*(\d+) (\d+) (\d+) (\d+) (\d+) (\d+)$")


def parse_attendance_pdf(text: str) -> tuple[list[tuple], list[str]]:
    """21대까지의 회기별 PDF → [(이름, 정당, 회의일수, 출석, 결석, 청가, 출장, 결석신고서)], 회의일들.

    엑셀과 달리 **그 회기 것만** 있고 누적이 없다. 줄은 '이름 정당 상태… 여섯 숫자'.
    동명이인 괄호 뒤에 정당이 붙어 나오는 줄이 있다: '이수진(비)더불어민주당 출석 …'.
    """
    days = sorted({f"{y}-{m}-{dd}" for y, m, dd in re.findall(r"(\d{4})년(\d{2})월(\d{2})일", text)})
    out = []
    for line in text.splitlines():
        line = line.strip()
        fused = re.match(r"^(\S+?\))(\S+)\s+(.*)$", line)
        if fused and not fused.group(2)[0].isdigit():
            line = f"{fused.group(1)} {fused.group(2)} {fused.group(3)}"
        m = ATT_LINE.match(line)
        if m:
            out.append((m.group(1), m.group(2), *(int(m.group(k)) for k in range(4, 10))))
    return out, days


def attendance_from_pdfs(c: httpx.Client, files: list[dict], age: int) -> tuple[list[tuple], int, str | None]:
    """그 대수 회기 PDF 를 전부 읽어 이름별로 더한다. 회기는 PDF 안 회의 날짜로 가른다
    (파일 목록에는 대수가 없다). 정당은 마지막으로 나온 회기 것을 쓴다."""
    import pdfplumber  # 무거워서 필요할 때만

    total: dict[str, list] = {}
    last_no, last_day = 0, None
    for f in files:
        if f.get("fileExt") != "pdf":
            continue
        r = c.get(ATT_DOWN, params={"infId": ATT_INF, "infSeq": 1, "fileSeq": f["fileSeq"]})
        r.raise_for_status()
        with pdfplumber.open(io.BytesIO(r.content)) as doc:
            text = "\n".join((pg.extract_text() or "") for pg in doc.pages)
        rows, days = parse_attendance_pdf(text)
        if not days or age_at(days[0]) != age:
            continue
        if not rows:
            raise RuntimeError(f"{f['viewFileNm']}: 출결 줄을 하나도 못 읽었습니다")
        for name, party, *v in rows:
            t = total.setdefault(name, [party, 0, 0, 0, 0, 0, 0])
            t[0] = party
            for k in range(6):
                t[k + 1] += v[k]
        last_no, last_day = max(last_no, f["no"]), max(filter(None, [last_day, days[-1]]))
        print(f"  {f['viewFileNm']}: {len(rows)}명, 회의 {len(days)}일", file=sys.stderr)
    return [(n, t[0], *t[1:]) for n, t in total.items()], last_no, last_day


def parse_attendance(rows: list[tuple]) -> tuple[list[tuple], str | None]:
    """출결 시트 → [(이름, 정당, 회의일수, 출석, 결석, 청가, 출장, 결석신고서)], 마지막 회의일.

    머리글: '구분 | 438회(임시) … | 총 계' 줄, 그 아래 '의원명 | 소속정당 | 1차 … | 회의일수 …',
    그 아래 회의 날짜 '(2026년08월26일)'. 누적은 '총 계' 칸부터 여섯 칸이다.
    """
    top = next(i for i, r in enumerate(rows) if r and "총 계" in [str(v).strip() for v in r if v])
    head, dates = rows[top + 1], rows[top + 2]
    ti = [str(v).strip() if v else "" for v in rows[top]].index("총 계")
    if tuple(str(v).strip() for v in head[ti:ti + 6]) != ATT_COLS:
        raise RuntimeError(f"출결 표 머리글이 바뀌었습니다: {head[ti:ti + 6]}")
    days = [ymd(v) for v in dates if v and re.search(r"\d{4}년", str(v))]
    out = []
    for r in rows[top + 3:]:
        if not r or not r[0]:
            continue
        out.append((str(r[0]).strip(), d(r[1]), *(int(r[ti + k] or 0) for k in range(6))))
    return out, max(days) if days else None


def ingest_attendance(cur, age: int) -> int:
    """그 대수 본회의 출결 누적.

    22대(제415회~)는 회기마다 엑셀이고 최신 파일에 대수 누적 '총 계' 가 있어 하나면 된다.
    21대까지는 PDF 이고 누적이 없어 그 대수 회기 PDF 를 전부 더한다(21대 34개).
    """
    import openpyxl  # 이 단계에서만 쓴다

    with httpx.Client(timeout=120, headers={"User-Agent": "nureongso/0.1"}) as c:
        files = attendance_files(c)
        f = max((x for x in files if x.get("fileExt") == "xlsx"), key=lambda x: x["no"])
        url = f"{ATT_DOWN}?infId={ATT_INF}&infSeq=1&fileSeq={f['fileSeq']}"
        r = c.get(url)
        r.raise_for_status()
        wb = openpyxl.load_workbook(io.BytesIO(r.content), data_only=True, read_only=True)
        ws = next((w for w in wb.worksheets if w.title.strip() == f"{age}대"), None)
        if ws is not None:
            rows, as_of = parse_attendance(list(ws.iter_rows(values_only=True)))
            session = f["no"]
            print(f"  {f['viewFileNm']} (마지막 회의 {as_of})", file=sys.stderr)
        else:
            rows, session, as_of = attendance_from_pdfs(c, files, age)
            url = ATT_PAGE  # 여러 파일을 더했으니 목록 화면으로 보낸다
            if not rows:
                raise RuntimeError(f"{age}대 출결 파일이 없습니다 (엑셀 시트: {wb.sheetnames})")

    # 단서가 있는 이름(한자, 괄호)부터 붙여야 같은 이름의 맨 한글 쪽이 '이미 붙은 사람
    # 제외' 로 갈린다 — 21대 김병욱·金炳旭, 이수진(비)·이수진.
    people = people_of(cur, age)
    used: set[str] = set()
    code_of: dict[str, str | None] = {}
    plain = lambda n: not (HANJA.match(n.replace(" ", "")) or "(" in n)
    for name, party, *_ in sorted(rows, key=lambda x: plain(x[0])):
        code_of[name] = match_person(people, name, party, used)
        if code_of[name]:
            used.add(code_of[name])
    bad = [x[0] for x in rows if x[2] != sum(x[3:8])]
    if bad:  # 회의일수 = 출석+결석+청가+출장+결석신고서 (실측 전원 성립)
        print(f"  ! 합이 안 맞는 줄 {len(bad)}: {', '.join(bad[:10])}", file=sys.stderr)
    miss = [n for n, c_ in code_of.items() if not c_]
    if miss:
        print(f"  ! 사람을 못 찾은 줄 {len(miss)}: {', '.join(miss)}", file=sys.stderr)

    cur.execute("delete from attendance where age = %s", (age,))
    cur.executemany(
        "insert into attendance (age, name, party, member_code, session_no, as_of, days,"
        " present, absent, leave, trip, absence_report, source_url)"
        " values (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)",
        [(age, n, p, code_of[n], session, as_of, *v, url) for n, p, *v in rows])
    return len(rows)


# --------------------------------------------------------------------------- run


STEPS = {
    "members": lambda cur, age: ingest_members(cur),
    "bills": lambda cur, age: ingest_bills(cur, age),
    "plenary": lambda cur, age: ingest_plenary(cur, age),
    "votes": lambda cur, age: ingest_votes(cur, age),
    # 의안 1건당 1회 호출이라 첫 실행이 오래 걸린다. 중간중간 커밋해서
    # 도중에 끊겨도 받은 만큼 남고, 다음 실행이 없는 것만 이어받는다.
    "summaries": lambda cur, age: ingest_summaries(cur, age),
    "sidejobs": lambda cur, age: ingest_sidejobs(cur, age),
    "attendance": lambda cur, age: ingest_attendance(cur, age),
    "trips": lambda cur, age: ingest_trips(cur, age),
    "research": lambda cur, age: ingest_research(cur, age),
}


def main():
    p = argparse.ArgumentParser()
    p.add_argument("step", choices=[*STEPS, "refresh", "all"])
    p.add_argument("--age", type=int, default=22)
    args = p.parse_args()

    dsn = os.environ.get("DATABASE_URL")
    if not dsn:
        sys.exit("DATABASE_URL 이 없습니다. .env 를 확인하세요.")

    steps = list(STEPS) if args.step == "all" else ([] if args.step == "refresh" else [args.step])

    failed: list[str] = []
    with psycopg.connect(dsn, autocommit=False) as conn:
        for name in steps:
            with conn.cursor() as cur:
                cur.execute("insert into ingest_run (source) values (%s) returning id",
                            (f"{name}:{args.age}",))
                run_id = cur.fetchone()[0]
            conn.commit()
            print(f"[{name}] 시작", file=sys.stderr)
            try:
                with conn.cursor() as cur:
                    n = STEPS[name](cur, args.age)
                conn.commit()
                with conn.cursor() as cur:
                    cur.execute("update ingest_run set finished_at=now(), rows=%s where id=%s",
                                (n, run_id))
                conn.commit()
                print(f"[{name}] 완료 {n} rows", file=sys.stderr)
            except Exception as e:
                conn.rollback()
                with conn.cursor() as cur:
                    cur.execute("update ingest_run set finished_at=now(), error=%s where id=%s",
                                (str(e)[:2000], run_id))
                conn.commit()
                # 한 단계가 막혔다고 나머지를 버리지 않는다. 국회 API 가 간헐적으로
                # 연결을 안 받는데, 첫 단계인 members 에서 걸리면 법안·표결·요약까지
                # 하루치를 통째로 건너뛰었다. 단계끼리는 서로 의존하지 않는다.
                # 실패는 ingest_run 에 남고 맨 끝에서 0 이 아닌 코드로 끝난다.
                print(f"[{name}] 실패: {str(e)[:200]}", file=sys.stderr)
                failed.append(name)

        if args.step in ("refresh", "all"):
            with conn.cursor() as cur:
                cur.execute("refresh materialized view concurrently member_stats")
                cur.execute("refresh materialized view concurrently member_party_line")
                # 매일 도는 수집이 단체장·교육감의 정당·지역구를 의원 시절로 되돌린
                # 적이 있다. 조용히 틀리면 아무도 모르니 매번 세어서 찍는다.
                cur.execute("""
                    with latest as (
                      select distinct on (member_code) member_code, office, party, district
                      from candidacy
                      where elected and member_code is not null and office is not null
                      order by member_code, election_id desc)
                    select count(*) from member m join latest l on l.member_code = m.code
                    where m.is_incumbent and l.office <> '국회의원'
                      and (m.party is distinct from l.party
                           or m.district is distinct from l.district)
                """)
                stale = cur.fetchone()[0]
            conn.commit()
            print("[refresh] member_stats · member_party_line 갱신", file=sys.stderr)
            if stale:
                print(f"[refresh] 경고: 현직 {stale}명의 정당·지역구가 당선 기록과"
                      " 다릅니다. `nec.py winners` 를 돌리세요.", file=sys.stderr)

    # 성공한 단계는 이미 커밋됐다. 그래도 조용히 넘어가면 안 된다 — 워크플로가
    # 초록불이면 며칠째 안 들어오는 데이터를 아무도 눈치채지 못한다.
    if failed:
        sys.exit(f"실패한 단계: {', '.join(failed)}")


if __name__ == "__main__":
    main()
